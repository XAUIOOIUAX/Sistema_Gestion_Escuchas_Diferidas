"""Exportación de un caso a Excel (.xlsx).

Reproduce el formato del libro original: una hoja TOTAL con todas las columnas
(incluida ESTADO) y una hoja INDICE con los abonados y su color aplicado como
relleno. Deja traza en LoteExportacion + LogAuditoria.
"""

from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import sesion
from app import repositorios as repo

# Encabezados de la hoja TOTAL (orden de columnas del libro original + ESTADO).
_ENCABEZADOS_TOTAL = [
    "ORDEN",
    "INTERESADO",
    "INTERLOCUTOR",
    "CD",
    "DIRECCION",
    "ORIGEN",
    "DESTINO",
    "CORRESPONDE A",
    "FECHA INICIO",
    "FECHA FIN",
    "ANTENAS",
    "LOCALIDAD",
    "PROVINCIA",
    "LATITUD",
    "LONGITUD",
    "AZIMUTH",
    "RADIO",
    "CONTEXTO",
    "ARCHIVO AUDIO",
    "ESTADO",
]

# Campo de Registro correspondiente a cada encabezado (None = derivado/constante).
_CAMPOS_TOTAL = [
    "orden",
    "interesado",
    "interlocutor",
    "cd",
    "direccion",
    "origen",
    "destino",
    "corresponde_a",
    "fecha_inicio_texto",
    "fecha_fin_texto",
    "antenas",
    "localidad",
    "provincia",
    "latitud",
    "longitud",
    "azimuth",
    "radio",
    "contexto",
    "archivo_audio",
    "estado",
]

# Columnas que deben forzarse a texto para que Excel no reinterprete (BUG fechas/números).
_COLS_TEXTO = {"FECHA INICIO", "FECHA FIN", "LATITUD", "LONGITUD", "AZIMUTH", "RADIO",
               "ORIGEN", "DESTINO"}

_HEADER_FILL = PatternFill("solid", fgColor="1E242C")
_HEADER_FONT = Font(color="E7E5DF", bold=True)


def _hex_openpyxl(color_hex: str | None) -> str | None:
    """Convierte '#RRGGBB' a 'RRGGBB' para openpyxl, o None si no hay color."""
    if not color_hex:
        return None
    return color_hex.lstrip("#").upper()


def exportar_caso_a_excel(
    con: sqlite3.Connection,
    caso_id: int,
    path_destino: str | Path,
    usuario: str | None = None,
) -> Path:
    """Genera el .xlsx del caso y devuelve la ruta del archivo."""
    usuario = sesion.o_analista(usuario)
    caso = repo.obtener_caso(con, caso_id)
    if caso is None:
        raise ValueError(f"Caso inexistente: {caso_id}")

    registros = repo.listar_registros(con, caso_id)
    # La hoja ÍNDICE es la leyenda de colores: se listan los abonados
    # intervenidos (cuyo color pinta las filas). Si el caso no tiene ninguno
    # marcado (importado antes de la detección), se listan todos como fallback.
    abonados = con.execute(
        "SELECT numero_normalizado, pertenece_a, color_hex, observaciones "
        "FROM Abonado WHERE caso_id = ? AND intervenido = 1 ORDER BY id",
        (caso_id,),
    ).fetchall()
    if not abonados:
        abonados = con.execute(
            "SELECT numero_normalizado, pertenece_a, color_hex, observaciones "
            "FROM Abonado WHERE caso_id = ? ORDER BY id",
            (caso_id,),
        ).fetchall()

    wb = Workbook()
    _hoja_total(wb, registros)
    _hoja_indice(wb, abonados)

    path_destino = Path(path_destino)
    path_destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path_destino))

    lote_id = repo.crear_lote_exportacion(con, caso_id, "excel", usuario)
    repo.log_auditoria(
        con,
        "exporto",
        f"Excel del caso '{caso['nombre']}' -> {path_destino} (lote #{lote_id})",
        caso_id,
        usuario,
    )
    con.commit()
    return path_destino


def _hoja_total(wb: Workbook, registros) -> None:
    ws = wb.active
    ws.title = "TOTAL"
    ws.append(_ENCABEZADOS_TOTAL)
    for celda in ws[1]:
        celda.fill = _HEADER_FILL
        celda.font = _HEADER_FONT
        celda.alignment = Alignment(horizontal="left")

    idx_texto = {
        i + 1 for i, h in enumerate(_ENCABEZADOS_TOTAL) if h in _COLS_TEXTO
    }

    for fila_n, r in enumerate(registros, start=2):
        valores = [r[campo] if campo in r.keys() else "" for campo in _CAMPOS_TOTAL]
        ws.append(valores)
        relleno = _hex_openpyxl(r["color_hex"])
        for col in range(1, len(_ENCABEZADOS_TOTAL) + 1):
            celda = ws.cell(row=fila_n, column=col)
            if col in idx_texto:
                celda.number_format = "@"
            if relleno:
                celda.fill = PatternFill("solid", fgColor=relleno)

    _autoancho(ws, _ENCABEZADOS_TOTAL)
    ws.freeze_panes = "A2"


def _hoja_indice(wb: Workbook, abonados) -> None:
    ws = wb.create_sheet("INDICE")
    encabezados = ["COLOR", "PERTENECE A", "NUMERO", "OBSERVACIONES"]
    ws.append(encabezados)
    for celda in ws[1]:
        celda.fill = _HEADER_FILL
        celda.font = _HEADER_FONT

    for fila_n, a in enumerate(abonados, start=2):
        ws.append(["", a["pertenece_a"] or "", a["numero_normalizado"], a["observaciones"] or ""])
        relleno = _hex_openpyxl(a["color_hex"])
        if relleno:
            ws.cell(row=fila_n, column=1).fill = PatternFill("solid", fgColor=relleno)
        ws.cell(row=fila_n, column=3).number_format = "@"

    _autoancho(ws, encabezados)
    ws.freeze_panes = "A2"


def _autoancho(ws, encabezados: list[str]) -> None:
    """Ajusta el ancho de columna al contenido más largo (cap razonable)."""
    for col_idx, _ in enumerate(encabezados, start=1):
        letra = get_column_letter(col_idx)
        largo = max(
            [len(str(encabezados[col_idx - 1]))]
            + [
                len(str(ws.cell(row=f, column=col_idx).value or ""))
                for f in range(2, ws.max_row + 1)
            ]
        )
        ws.column_dimensions[letra].width = min(max(largo + 2, 8), 48)
